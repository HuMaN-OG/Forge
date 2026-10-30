from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
import os


# Local storage directory
LOCAL_APP_DATA = Path(
    os.environ.get("LOCALAPPDATA", Path.home() / "AppData" / "Local")
)

OUTPUT_DIR = LOCAL_APP_DATA / "ForgeAI" / "outputs"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def create_student_excel():

    file_path = OUTPUT_DIR / "students.xlsx"

    # Sample student data
    students = [
        ["Roll Number", "Name", "Marks"],
        [101, "Rahul", 85],
        [102, "Priya", 92],
        [103, "Aman", 78],
        [104, "Sneha", 88],
        [105, "Karan", 95]
    ]

    # Create workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"

    # Insert data
    for row in students:
        sheet.append(row)

    # Style header
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            start_color="2563EB",
            end_color="2563EB",
            fill_type="solid"
        )
        cell.alignment = Alignment(horizontal="center")

    # Adjust column widths
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 15

    # Save workbook
    workbook.save(file_path)

    # Verify saved workbook
    verified_workbook = load_workbook(file_path)
    verified_sheet = verified_workbook["Students"]

    assert verified_sheet.max_row == 6
    assert verified_sheet["B2"].value == "Rahul"

    verified_workbook.close()

    return {
        "status": "success",
        "message": "Excel file created and verified",
        "file_path": str(file_path),
        "total_students": 5
    }


if __name__ == "__main__":
    result = create_student_excel()
    print(result)